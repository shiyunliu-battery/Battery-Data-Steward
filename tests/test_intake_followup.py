"""Regression tests for realistic export formatting and semantic ambiguity."""

import math

import polars as pl
import pytest
from openpyxl import Workbook

import bds
from battery_data_standard.adapters.normalization import _parse_number, _parse_unix_time
from battery_data_standard.cli import main
from battery_data_standard.diagnostics import explain
from battery_data_standard.exceptions import ConversionError, UnsupportedFormatError
from battery_data_standard.quality import _repeated_step_segments
from battery_data_standard.reporting import _report_statement, _time_sampling_rows
from battery_data_standard.time_sampling import apply_time_sampling_policy


@pytest.mark.parametrize("status", ["Dynamic", "Delay", "DCIR", "Control"])
def test_unknown_status_does_not_change_current(status, tmp_path):
    source = tmp_path / "drive.csv"
    source.write_text(f"Time,Voltage,Current,Status\n0,4,-1,{status}\n1,4,1,{status}\n")
    frame, report = bds.read_with_report(source, cycler="generic")
    assert frame["current_a"].to_list() == [-1, 1]
    assert any("preserved" in notice for notice in report.warnings)


def test_pec_formatted_hours_are_not_scaled_twice(tmp_path):
    source = tmp_path / "pec.csv"
    source.write_text(
        "Total Time (Hours in hh:mm:ss.xxx),Voltage (V),Current (A)\n0:00:00,4,1\n1:00:00,4,1\n"
    )
    frame = bds.read(source, cycler="pec")
    assert frame["test_time_s"].to_list() == [0, 3600]


@pytest.mark.parametrize("encoding", ["cp1252", "utf-8"])
def test_microamp_units_survive_decoding(encoding, tmp_path):
    source = tmp_path / "micro.csv"
    source.write_bytes("Time,Voltage,Current (µA)\n0,4,500\n1,4,500\n".encode(encoding))
    frame = bds.read(source, cycler="generic")
    assert frame["current_a"].to_list() == pytest.approx([0.0005, 0.0005])


def test_chinese_neware_headers_statuses_and_capacity(tmp_path):
    source = tmp_path / "neware.csv"
    source.write_bytes(
        "测试时间(s),电压(V),电流(A),循环号,工步号,容量(Ah),工步类型\n0,4,1,1,1,0.1,充电\n1,4,1,1,2,0.2,放电\n".encode(
            "gb18030"
        )
    )
    frame, report = bds.read_with_report(source, cycler="auto")
    assert report.cycler == "neware"
    assert frame["cycle_index"].to_list() == [1, 1]
    assert frame["step_index"].to_list() == [1, 2]
    assert frame["current_a"].to_list() == [1, -1]
    assert frame["charge_capacity_ah"].to_list() == [0.1, None]
    assert frame["discharge_capacity_ah"].to_list() == [None, 0.2]


def test_dotted_dates_are_dayfirst_and_slash_dates_can_be_configured(tmp_path):
    dotted = _parse_unix_time(["01.02.2023 00:00:00", "02.02.2023 00:00:00"])
    assert dotted[1] - dotted[0] == 86400
    source = tmp_path / "dates.csv"
    source.write_text("DateTime,Voltage,Current\n01/02/2023 00:00:00,4,1\n02/02/2023 00:00:00,4,1\n")
    frame = bds.read(source, cycler="generic", profile={"date_order": "dayfirst"})
    assert frame["test_time_s"].to_list() == [0, 86400]
    _, report = bds.read_with_report(source, cycler="generic")
    assert any("month-first" in value for value in report.warnings)


def test_biologic_flag_is_not_a_cycle_and_undirected_capacity_is_not_charge(tmp_path):
    source = tmp_path / "biologic.csv"
    source.write_text("time/s,Ewe/V,I/mA,Ns changes,Capacity/mA.h\n0,4,-100,0,2\n1,4,-100,1,3\n")
    frame, report = bds.read_with_report(source, cycler="biologic")
    assert "cycle_index" not in frame.columns
    assert "charge_capacity_ah" not in frame.columns
    assert frame["source_capacity_ah"].to_list() == [0.002, 0.003]
    assert "Ns changes" in report.unmapped_columns


@pytest.mark.parametrize(
    "cycler, header, factor", [("neware", "Capacity(mAh)", 0.001), ("maccor", "Amp-hr", 1.0)]
)
def test_undirected_vendor_capacity_is_split_by_explicit_status(cycler, header, factor, tmp_path):
    source = tmp_path / "capacity.csv"
    source.write_text(f"Test Time,Voltage,Current,Status,{header},Watt-hr\n0,4,1,C,2,8\n1,4,1,D,1,4\n")
    frame = bds.read(source, cycler=cycler)
    assert frame["charge_capacity_ah"].to_list() == [2 * factor, None]
    assert frame["discharge_capacity_ah"].to_list() == [None, factor]
    assert frame["discharge_energy_wh"].to_list() == [None, 4]


def test_source_power_is_preserved_with_sign_warning(tmp_path):
    source = tmp_path / "power.csv"
    source.write_text("Time,Voltage,Current,Status,Power(W)\n0,4,1,D,4\n1,4,1,D,4\n")
    frame, report = bds.read_with_report(source, cycler="generic")
    assert frame["current_a"].to_list() == [-1, -1]
    assert frame["power_w"].to_list() == [4, 4]
    assert any("Source power" in notice for notice in report.warnings)


def test_row_limit_counts_only_actual_insertions_and_preserves_null_dtype():
    frame = pl.DataFrame(
        {
            "test_time_s": [0.0, 2.0, 1002.0],
            "current_a": [1.0, 1.0, -1.0],
            "step_index": [1, 1, 2],
            "raw:empty": [None] * 3,
        }
    )
    result = apply_time_sampling_policy(frame, policy="repair", expected_interval_s=1, max_inserted_rows=1)
    assert result.data["test_time_s"].to_list() == [0, 1, 2, 1002]
    assert result.data.schema == frame.schema
    assert result.metadata["boundary_skipped_points"] == 999
    assert result.metadata["inserted_rows"] == 1
    payload = {"status": "ok", "time_sampling": result.metadata}
    assert "interpolated" in _report_statement(payload)
    assert "999" in _report_statement(payload)
    assert ["Boundary points left unfilled", 999] in _time_sampling_rows(result.metadata)


def test_large_boundary_gap_does_not_allocate_missing_grid():
    frame = pl.DataFrame({"test_time_s": [0.0, 2.0, 1e12], "step_index": [1, 1, 2]})
    result = apply_time_sampling_policy(frame, policy="repair", expected_interval_s=1, max_inserted_rows=1)
    assert result.data.height == 4
    assert len(result.metadata["gaps"][1]["missing_times_s"]) == 20


def test_explain_and_convert_have_matching_default_sampling(tmp_path):
    source = tmp_path / "gap.csv"
    source.write_text("Time,Voltage,Current\n0,4,1\n1,4,1\n2,4,1\n4,4,1\n")
    assert explain(source).time_sampling["status"] == "gaps-detected"
    assert bds.read(source).height == 4


@pytest.mark.parametrize("padding", ["\n\n\n", "\n,,\n", "\nFrequency_Hz,Zreal,Zimag\n"])
def test_eis_skips_only_blank_or_repeated_header_rows_with_report(padding, tmp_path):
    source = tmp_path / "eis.csv"
    source.write_text("Frequency_Hz,Zreal,Zimag\n1000,.1,-.01" + padding + "100,.2,-.02\n")
    output = tmp_path / "out.csv"
    report = bds.convert_eis(source, output)
    assert report.rows == 2
    assert any(issue.code == "eis-normalization-notice" for issue in report.issues)


def test_phase_without_unit_warns_and_accepts_explicit_radians(tmp_path):
    source = tmp_path / "phase.csv"
    source.write_text(f"Frequency_Hz,Zmod,Phase\n1000,1,{math.radians(-5)}\n")
    with pytest.warns(UserWarning, match="unspecified"):
        bds.read_eis(source)
    frame = bds.read_eis(source, phase_unit="rad")
    assert frame["Phase_exp_deg"][0] == pytest.approx(-5)
    report = bds.convert_eis(source, tmp_path / "out.csv")
    assert any("unspecified" in issue.message for issue in report.issues)


def test_multisheet_batch_keeps_successes_errors_and_later_sheets(tmp_path):
    source = tmp_path / "mixed.xlsx"
    workbook = Workbook()
    for index, name in enumerate(["good", "bad", "later"]):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        sheet.append(["Frequency_Hz", "Zreal", "Zimag"])
        sheet.append([1000, None if name == "bad" else 0.1, -0.01])
    workbook.save(source)
    records = bds.batch_convert(source, tmp_path / "out", excel_sheets="all")
    assert [(record["sheet_name"], record["status"]) for record in records] == [
        ("good", "ok"),
        ("bad", "error"),
        ("later", "ok"),
    ]
    assert (tmp_path / "out" / "mixed_good.eis.csv").exists()
    assert (tmp_path / "out" / "mixed_later.eis.csv").exists()
    with pytest.raises(ConversionError):
        bds.batch_convert(source, tmp_path / "fast", excel_sheets="all", fail_fast=True)


def test_multiple_arbin_channels_require_explicit_selection(tmp_path):
    source = tmp_path / "channels.xlsx"
    workbook = Workbook()
    for index in range(2):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = f"Channel_{index + 1}"
        sheet.append(["Test_Time(s)", "Voltage(V)", "Current(A)"])
        sheet.append([0, 4, 1])
        sheet.append([1, 4, 1])
    workbook.save(source)
    with pytest.raises(UnsupportedFormatError, match="Multiple channel"):
        bds.read(source, cycler="arbin")
    assert bds.read(source, cycler="arbin", sheet="Channel_1").height == 2
    records = bds.batch_convert(source, tmp_path / "out", excel_sheets="all")
    assert len(records) == 2 and all(record["status"] == "ok" for record in records)


def test_ambiguous_comma_numbers_are_rejected():
    with pytest.raises(ConversionError, match="Ambiguous numeric"):
        _parse_number("1,000")
    assert _parse_number("1,25") == 1.25


def test_pybamm_default_uses_discharge_positive_and_explicit_override(tmp_path):
    source = tmp_path / "discharge.csv"
    source.write_text("Time,Voltage,Current,Status\n0,4,1,D\n1,4,1,D\n")
    output = tmp_path / "profile.csv"
    bds.convert(source, output, target="pybamm")
    assert pl.read_csv(output)["current_a"].to_list() == [1, 1]
    assert main(["convert", str(source), str(output), "--target", "pybamm"]) == 0
    assert pl.read_csv(output)["current_a"].to_list() == [1, 1]
    bds.convert(source, output, target="pybamm", current_sign="charge-positive")
    assert pl.read_csv(output)["current_a"].to_list() == [-1, -1]


def test_cycle_capacity_requires_reset_scope_and_sums_step_reset_values():
    frame = pl.DataFrame(
        {"cycle_index": [1] * 4, "step_index": [1, 1, 2, 2], "charge_capacity_ah": [0.0, 2.0, 0.0, 1.0]}
    )
    with pytest.warns(UserWarning, match="reset scope"):
        summary = bds.summarize_cycles(frame)
    assert summary["Charging Capacity Recorded Maximum / Ah"][0] == 2
    assert bds.summarize_cycles(frame, capacity_reset="step")["Charging Capacity End / Ah"][0] == 3
    assert bds.summarize_cycles(frame, capacity_reset="cycle")["Charging Capacity End / Ah"][0] == 2


def test_repeated_step_segments_matches_interleaved_cycle_semantics():
    assert _repeated_step_segments([1, 2, 1, 2, 1], [1, 1, 2, 2, 1]) == {
        "repeated_step_segments": 1,
        "repeated_step_examples": [{"cycle": 1, "step": 1, "row": 4}],
    }


@pytest.mark.parametrize(
    "duration, expected", [("25:00:00", 90000), ("1-01:00:00", 90000), ("1 days 01:00:00", 90000)]
)
def test_long_durations_and_generic_units(duration, expected, tmp_path):
    source = tmp_path / "durations.csv"
    source.write_text(f"Time,Voltage (mV),Current\n00:00:00,4000,1\n{duration},4000,1\n")
    frame = bds.read(source, cycler="generic")
    assert frame["test_time_s"].to_list() == [0, expected]
    assert frame["voltage_v"].to_list() == [4, 4]


def test_formatted_hour_header_and_excel_long_duration(tmp_path):
    import datetime

    source = tmp_path / "hours.csv"
    source.write_text("Time [h]:mm:ss,Voltage,Current\n00:00:00,4,1\n25:00:00,4,1\n")
    assert bds.read(source, cycler="generic")["test_time_s"].to_list() == [0, 90000]
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Time", "Voltage", "Current"])
    sheet.append([datetime.timedelta(), 4, 1])
    sheet.append([datetime.timedelta(hours=25), 4, 1])
    sheet["A2"].number_format = "[h]:mm:ss"
    sheet["A3"].number_format = "[h]:mm:ss"
    path = tmp_path / "hours.xlsx"
    workbook.save(path)
    assert bds.read(path, cycler="generic")["test_time_s"].to_list() == [0, 90000]


def test_generic_minute_millivolt_and_milliohm_units(tmp_path):
    source = tmp_path / "units.csv"
    source.write_text(
        "Time [min],Voltage (mV),Current,Internal Resistance (mΩ)\n0,4000,1,10\n1,4000,1,10\n",
        encoding="utf-8",
    )
    frame = bds.read(source, cycler="generic")
    assert frame["test_time_s"].to_list() == [0, 60]
    assert frame["voltage_v"].to_list() == [4, 4]
    assert frame["internal_resistance_ohm"].to_list() == [0.01, 0.01]


def test_german_headers_do_not_collapse_into_empty_ascii_slugs(tmp_path):
    source = tmp_path / "german.csv"
    source.write_text("Programmdauer,Spannung(V),Strom(A)\n0,4,1\n1,4,1\n")
    assert bds.read(source, cycler="generic").height == 2


def test_integer_input_interpolation_retains_fractional_values():
    frame = pl.DataFrame({"test_time_s": [0, 1, 3], "current_a": [0, 0, 1]})
    result = apply_time_sampling_policy(frame, policy="repair", expected_interval_s=1)
    assert result.data["current_a"].to_list() == [0, 0, 0.5, 1]


def test_nonempty_invalid_eis_and_zero_frequency_still_fail(tmp_path):
    source = tmp_path / "bad.csv"
    for row in ["1000,NaN,-.01", "0,.1,-.01"]:
        source.write_text("Frequency_Hz,Zreal,Zimag\n1000,.1,-.01\n" + row)
        with pytest.raises(ConversionError):
            bds.convert_eis(source, tmp_path / "out.csv")


def test_wide_unknown_layout_is_rejected_instead_of_guessed(tmp_path):
    source = tmp_path / "unknown.csv"
    source.write_text(",".join(str(i) for i in range(10)) + "\n" + ",".join(str(i + 1) for i in range(10)))
    with pytest.raises(UnsupportedFormatError, match="Headerless wide tables"):
        bds.read(source, cycler="generic")


@pytest.mark.parametrize("separator, expected", [(",", 1.0), (".", 1000.0)])
def test_explicit_decimal_separator_resolves_three_digit_comma_values(separator, expected, tmp_path):
    source = tmp_path / "locale.csv"
    source.write_text("Time;Voltage;Current\n0;4;1,000\n1;4;1,000\n")
    frame = bds.read(source, cycler="generic", profile={"decimal_separator": separator})
    assert frame["current_a"].to_list() == [expected, expected]
